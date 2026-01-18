import csv
from typing import List, Dict, Any
from datetime import timedelta
import openpyxl
from openpyxl.utils import get_column_letter

class Exporter:
    @staticmethod
    def _format_duration(seconds: int) -> str:
        if seconds is None:
            return ""
        m, s = divmod(seconds, 60)
        h, m = divmod(m, 60)
        return f"{h:02}:{m:02}:{s:02}"

    @staticmethod
    def export_csv(filepath: str, sessions: List[Dict[str, Any]], fields: List[str] = None):
        if not sessions:
            return

        # Default fields if not provided
        # Added Project name to export
        if not fields:
            fields = ['id', 'date', 'project_name', 'start_ts', 'end_ts', 'duration_sec', 'comment', 'mode']
        
        # Map internal keys to display headers if needed
        # We handle missing keys gracefully
        
        try:
            with open(filepath, 'w', newline='', encoding='utf-8') as f:
                writer = csv.DictWriter(f, fieldnames=fields)
                writer.writeheader()
                
                for session in sessions:
                    # Filter only requested fields
                    row = {k: session.get(k) for k in fields}
                    writer.writerow(row)
        except Exception as e:
            raise IOError(f"Failed to export CSV: {e}")

    @staticmethod
    def export_excel(filepath: str, sessions: List[Dict[str, Any]], fields: List[str] = None):
        if not sessions:
            return

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Work Sessions"
        
        # Headers including Project
        headers = ['Date', 'Project', 'Start Time', 'End Time', 'Duration (HH:MM:SS)', 'Duration (Minutes)', 'Comment', 'Mode']
        ws.append(headers)
        
        # Style headers
        for cell in ws[1]:
            cell.font = openpyxl.styles.Font(bold=True)
            
        for session in sessions:
            # Derived values
            date = session.get('date')
            project = session.get('project_name')
            start = session.get('start_ts')
            end = session.get('end_ts')
            dur_sec = session.get('duration_sec')
            comment = session.get('comment')
            mode = session.get('mode')
            
            dur_fmt = Exporter._format_duration(dur_sec) if dur_sec is not None else ""
            dur_min = round(dur_sec / 60, 2) if dur_sec is not None else 0
            
            ws.append([date, project, start, end, dur_fmt, dur_min, comment, mode])
            
        # Auto-width
        for column_cells in ws.columns:
            length = max(len(str(cell.value) or "") for cell in column_cells)
            ws.column_dimensions[get_column_letter(column_cells[0].column)].width = length + 2
            
        try:
            wb.save(filepath)
        except Exception as e:
            raise IOError(f"Failed to export Excel: {e}")
